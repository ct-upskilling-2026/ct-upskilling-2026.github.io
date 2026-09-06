"""Generate the ABC Tutoring customer deck from presentation/SLIDE-SPEC.md.

    tools/../<venv>/bin/python tools/build_deck.py

Produces presentation/ABCTutoringPresentation.pptx (16:9). Export to PDF from
PowerPoint with File > Export > PDF.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches as In, Pt, Emu
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

ROOT = Path(__file__).resolve().parent.parent
IMG  = ROOT / "presentation" / "img"

CREAM  = C(0xFD, 0xFB, 0xF7); WHITE = C(0xFF, 0xFF, 0xFF)
TEAL   = C(0x2F, 0x6F, 0x62); TEAL_D = C(0x24, 0x54, 0x49)
TEAL_L = C(0x9C, 0xC4, 0xBB)
AMBER  = C(0xF2, 0xA6, 0x5A); INK   = C(0x2B, 0x2B, 0x2B)
MUTED  = C(0x5F, 0x5B, 0x54); BORDER = C(0xEA, 0xE3, 0xD8)
FONT   = "Avenir Next"

prs = Presentation()
prs.slide_width, prs.slide_height = In(13.333), In(7.5)
W = 13.333


def slide(notes=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = CREAM
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    """runs = [(string, size, bold, color), ...] — one paragraph each."""
    tb = s.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (t, size, bold, col) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run(); r.text = t
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = col
    return tb


def title(s, t):
    text(s, 0.62, 0.42, 11.5, 0.7, [(t, 33, True, TEAL)])


def card(s, x, y, w, h, fill=WHITE, line=BORDER, lw=1.0, radius=0.06):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, In(x), In(y), In(w), In(h))
    sh.adjustments[0] = radius
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def badge(s, x, y, d, label, fill, txt_col=WHITE, size=15):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, In(x), In(y), In(d), In(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background(); sh.shadow.inherit = False
    tf = sh.text_frame; tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size, r.font.bold, r.font.name = Pt(size), True, FONT
    r.font.color.rgb = txt_col
    return sh


def framed_image(s, path, cx, top, h):
    """Place image scaled to height h, centered on cx, inside a white frame."""
    from struct import unpack
    d = open(path, "rb").read(33)
    iw, ih = unpack(">II", d[16:24])
    w = h * iw / ih
    pad = 0.075
    card(s, cx - w / 2 - pad, top - pad, w + 2 * pad, h + 2 * pad, radius=0.04)
    s.shapes.add_picture(str(path), In(cx - w / 2), In(top), In(w), In(h))
    return w


def style_chart(ch, size=11):
    ch.font.size, ch.font.name = Pt(size), FONT
    ch.font.color.rgb = MUTED
    try:
        ch.has_title = False
    except Exception:
        pass


# ─────────────────────────── Slide 1 ───────────────────────────
s = slide("A listening check, not a status report — pause here and let Dana confirm each "
          "line matches what she asked for before moving on.")
title(s, "What I heard from you")
lines = [
    "Right now everything runs through your phone and a Google Sheet",
    "You want a simple site: a home page, a page to meet the tutors, and a way to book",
    "Each tutor should show their photo, subjects, grade levels, rate, and real availability",
    "Booking should just ask for the parent's name and email, and the student's name, grade, and subject",
    "The calendar has to show each tutor's actual free times — not a generic 9-to-5",
    "It needs to feel clean and friendly on a phone — not dark, not corporate, not a big test-prep chain",
]
y = 1.62
for ln in lines:
    badge(s, 0.72, y - 0.02, 0.38, "✓", TEAL, WHITE, 15)
    text(s, 1.32, y, 11.2, 0.5, [(ln, 17, False, INK)], anchor=MSO_ANCHOR.TOP)
    y += 0.86

# ─────────────────────────── Slide 2 ───────────────────────────
s = slide("Live screenshots from the real site at https://ct-upskilling-2026.github.io, shown "
          "at phone width because that is how most parents will actually see it.")
title(s, "Your website")
caps = [("index.png",     "Home — a warm welcome and a clear next step"),
        ("tutors.png",    "Meet the tutors — photo, subjects, rate, and open times"),
        ("book.png",      "Book a session — one short form, done in about a minute")]
for i, (fn, cap) in enumerate(caps):
    cx = 2.6 + i * 4.07
    framed_image(s, IMG / fn, cx, 1.42, 4.72)
    text(s, cx - 1.85, 6.36, 3.7, 0.7, [(cap, 13, False, MUTED)], align=PP_ALIGN.CENTER, spacing=1.15)

# ─────────────────────────── Slide 3 ───────────────────────────
s = slide("Emphasize 'the slot closes automatically' — that single behavior is what eliminates "
          "the double-booking risk she has today with the Google Sheet and her phone.")
title(s, "How a parent books")
steps = [
    ("Pick a tutor", "See their subjects, rate, and photo"),
    ("Pick an open time slot", "Only real, available times are ever shown"),
    ("Fill in one short form", "Parent and student details — nothing more"),
    ("Get an instant confirmation", "That slot closes automatically, so nobody double-books it"),
]
y = 1.55
for i, (h, sub) in enumerate(steps, 1):
    badge(s, 0.72, y, 0.52, str(i), AMBER, C(0x3A, 0x2A, 0x14), 18)
    text(s, 1.46, y - 0.02, 6.6, 0.42, [(h, 18, True, INK)])
    text(s, 1.46, y + 0.36, 6.6, 0.42, [(sub, 14.5, False, MUTED)])
    y += 1.06
card(s, 0.72, 6.06, 7.35, 0.82, fill=C(0xEF, 0xF5, 0xF3), line=C(0xCF, 0xE3, 0xDD), radius=0.14)
text(s, 0.95, 6.28, 6.9, 0.45, [("Start to finish: about a minute. No phone tag.", 19, True, TEAL_D)])
framed_image(s, IMG / "book.png",      10.75, 1.5,  2.5)
framed_image(s, IMG / "confirmed.png", 10.75, 4.35, 2.5)

# ─────────────────────────── Slide 4 ───────────────────────────
s = slide("These are simulated demo numbers from two weeks of test traffic (340 visitors), not real "
          "customers — the measurement is proven and ready to read real traffic once the site is "
          "promoted. Reiterate verbally that this is a recommendation: Dana was asked twice what she "
          "wanted to know about visitors and both times described booking-form fields instead, which "
          "is exactly why this slide is framed as 'here is what I'd suggest'.")
title(s, "What you can now see about your visitors")
text(s, 0.62, 1.06, 11.6, 0.35,
     [("A suggestion from me: here are the four things I'd watch, and what each one is worth to you.",
       14, False, MUTED)])

card(s, 0.62, 1.52, 12.1, 0.92, fill=WHITE, line=AMBER, lw=2.25, radius=0.1)
text(s, 0.92, 1.66, 11.5, 0.4,
     [("44 parents came looking for Algebra II help. Only 4 could book. 10 found no open times.",
       19, True, INK)])
text(s, 0.92, 2.03, 11.5, 0.34,
     [("One tutor covers Algebra II with 2 hours a week. Compare Elementary Reading, which books at "
       "41% because that tutor has hours to spare.", 12.5, False, MUTED)])

# Chart A — looked vs booked, ascending so the best rate lands on top
cd = CategoryChartData()
cd.categories = ["Algebra II", "Science", "Elementary Math", "Algebra I", "Pre-Algebra", "Elem. Reading"]
cd.add_series("Looked", (44, 24, 48, 45, 27, 29))
cd.add_series("Booked", (4, 4, 10, 12, 10, 12))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, In(0.62), In(2.62), In(7.15), In(4.35), cd)
ch = gf.chart; style_chart(ch, 11)
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.plots[0].series[0].format.fill.solid(); ch.plots[0].series[0].format.fill.fore_color.rgb = TEAL_L
ch.plots[0].series[1].format.fill.solid(); ch.plots[0].series[1].format.fill.fore_color.rgb = TEAL_D
ch.plots[0].gap_width = 60
ch.plots[0].has_data_labels = True
dl = ch.plots[0].data_labels
dl.font.size = Pt(9.5); dl.font.name = FONT; dl.font.color.rgb = MUTED
dl.position = XL_LABEL_POSITION.OUTSIDE_END
ch.value_axis.has_major_gridlines = False
ch.value_axis.visible = False
text(s, 0.62, 2.5, 7.0, 0.3, [("Looked vs. booked, by subject", 13, True, INK)])

# Chart B — the drop-off, reversed so the first stage sits on top
cd = CategoryChartData()
cd.categories = ["Booked", "Started the form", "Picked a time",
                 "Opened a tutor", "Viewed the tutors", "Visited home page"]
cd.add_series("Parents", (52, 67, 81, 138, 217, 340))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, In(8.05), In(2.62), In(4.68), In(2.12), cd)
ch = gf.chart; style_chart(ch, 9.5)
ch.has_legend = False
ch.plots[0].series[0].format.fill.solid(); ch.plots[0].series[0].format.fill.fore_color.rgb = TEAL
ch.plots[0].gap_width = 45
ch.plots[0].has_data_labels = True
dl = ch.plots[0].data_labels; dl.font.size = Pt(9); dl.font.name = FONT; dl.font.color.rgb = MUTED
dl.position = XL_LABEL_POSITION.OUTSIDE_END
ch.value_axis.has_major_gridlines = False; ch.value_axis.visible = False
text(s, 8.05, 2.5, 4.6, 0.3, [("From visit to booking", 13, True, INK)])

# Chart C — how parents found you (bars label themselves; a donut hides the numbers)
cd = CategoryChartData()
cd.categories = ["Typed it in", "Flyer", "Word of mouth", "Facebook"]
cd.add_series("Parents", (43, 58, 81, 158))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, In(8.05), In(5.0), In(4.68), In(1.55), cd)
ch = gf.chart; style_chart(ch, 9.5)
ch.has_legend = False
ch.plots[0].series[0].format.fill.solid(); ch.plots[0].series[0].format.fill.fore_color.rgb = AMBER
ch.plots[0].gap_width = 45
ch.plots[0].has_data_labels = True
dl = ch.plots[0].data_labels
dl.font.size = Pt(9.5); dl.font.name = FONT; dl.font.color.rgb = MUTED
dl.position = XL_LABEL_POSITION.OUTSIDE_END
ch.value_axis.has_major_gridlines = False; ch.value_axis.visible = False
text(s, 8.05, 4.86, 4.6, 0.3, [("How parents found you", 13, True, INK)])

# The device split reads better as a sentence than as a two-slice donut
card(s, 8.05, 6.6, 4.68, 0.62, fill=C(0xEF, 0xF5, 0xF3), line=C(0xCF, 0xE3, 0xDD), radius=0.16)
text(s, 8.25, 6.6, 4.3, 0.62,
     [("69% of parents visited on a phone (233 of 340)", 12.5, True, TEAL_D)],
     anchor=MSO_ANCHOR.MIDDLE)

# ─────────────────────────── Slide 5 ───────────────────────────
s = slide("Be explicit and warm that sample tutors were used because real tutor details were never "
          "provided — this is not a gap in the work, it is the one open input still needed. Frame "
          "phase 2 as sequenced, not forgotten.")
title(s, "What I need from you")
card(s, 0.62, 1.45, 5.9, 5.05, radius=0.05)
text(s, 0.95, 1.78, 5.3, 0.4, [("To replace the sample tutors", 19, True, TEAL)])
y = 2.42
for b in ["Real names, photos, subjects, and grade levels for your 6 tutors",
          "Each tutor's hourly rate",
          "Each tutor's real weekly availability (the times shown now are placeholders)"]:
    badge(s, 0.98, y + 0.03, 0.16, "", AMBER)
    text(s, 1.32, y, 4.9, 0.85, [(b, 15.5, False, INK)], spacing=1.15)
    y += 1.0
card(s, 0.95, 5.35, 5.24, 0.92, fill=C(0xF8, 0xF5, 0xEF), line=BORDER, radius=0.12)
text(s, 1.18, 5.54, 4.8, 0.6,
     [("All of this lives in a single file — once you send it over, it's a quick swap. "
       "No rebuilding the site.", 13, False, MUTED)], spacing=1.15)

card(s, 6.82, 1.45, 5.9, 5.05, radius=0.05)
text(s, 7.15, 1.78, 5.3, 0.4, [("Coming next (not built yet)", 19, True, TEAL)])
y = 2.42
for b in ["A simple screen where you can add or remove tutors yourself",
          "A shared calendar so bookings stay in sync across every device",
          "Automatic email confirmations to parents",
          "Invoicing"]:
    badge(s, 7.18, y + 0.03, 0.16, "", TEAL_L)
    text(s, 7.52, y, 4.9, 0.85, [(b, 15.5, False, INK)], spacing=1.15)
    y += 0.92

out = ROOT / "presentation" / "ABCTutoringPresentation.pptx"
prs.save(out)
print("wrote", out, "—", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
